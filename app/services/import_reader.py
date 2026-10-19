import csv
import json
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.errors import AppError

MAX_IMPORT_BYTES = 10 * 1024 * 1024


def _rows_from_csv(content: bytes) -> list[tuple[int, dict[str, Any]]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise AppError(400, "invalid_import_encoding", "CSV files must use UTF-8") from exc
    return [(index, dict(row)) for index, row in enumerate(csv.DictReader(StringIO(text)), 2)]


def _rows_from_json(content: bytes) -> list[tuple[int, dict[str, Any]]]:
    try:
        payload = json.loads(content.decode("utf-8"))
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise AppError(400, "invalid_json_import", "JSON import content is invalid") from exc
    if isinstance(payload, dict):
        payload = payload.get("items")
    if not isinstance(payload, list) or not all(isinstance(item, dict) for item in payload):
        raise AppError(400, "invalid_json_import", "JSON import must be an array of objects")
    return [(index, item) for index, item in enumerate(payload, 1)]


def _rows_from_excel(content: bytes) -> list[tuple[int, dict[str, Any]]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(iterator)]
        return [
            (index, dict(zip(headers, values, strict=False)))
            for index, values in enumerate(iterator, 2)
        ]
    except (OSError, ValueError, StopIteration) as exc:
        raise AppError(400, "invalid_excel_import", "Excel import content is invalid") from exc


def read_import_rows(filename: str, content: bytes) -> list[tuple[int, dict[str, Any]]]:
    if len(content) > MAX_IMPORT_BYTES:
        raise AppError(413, "import_too_large", "Import file exceeds the 10 MB limit")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return _rows_from_csv(content)
    if suffix == ".json":
        return _rows_from_json(content)
    if suffix == ".xlsx":
        return _rows_from_excel(content)
    raise AppError(415, "unsupported_import_format", "Supported formats are CSV, JSON, and XLSX")
