from io import BytesIO

from docx import Document


def test_markdown_ingestion_hybrid_search_and_citations(client):
    document = (
        b"# Infusion pump maintenance\n\n"
        b"Occlusion alarms may be caused by a bent tube or a closed clamp.\n\n"
        b"Inspect the tube path and clamp before replacing components.\n"
    )

    ingestion = client.post(
        "/api/v1/knowledge/documents",
        data={"title": "Infusion Pump Guide"},
        files={"file": ("guide.md", document, "text/markdown")},
    )
    assert ingestion.status_code == 202
    assert ingestion.json()["status"] == "succeeded"

    response = client.post(
        "/api/v1/knowledge/search", json={"query": "occlusion alarm clamp", "top_k": 3}
    )
    assert response.status_code == 200
    hits = response.json()
    assert len(hits) >= 1
    assert hits[0]["title"] == "Infusion Pump Guide"
    assert "clamp" in hits[0]["content"].lower()
    assert hits[0]["citation"].startswith("[Infusion Pump Guide#")
    assert {"dense_score", "keyword_score", "fusion_score", "rerank_score"} <= set(
        hits[0]["scores"]
    )


def test_docx_ingestion_and_mock_question_answer(client):
    doc = Document()
    doc.add_heading("Patient monitor troubleshooting", level=1)
    doc.add_paragraph("If the ECG lead-off warning appears, inspect electrode contact first.")
    buffer = BytesIO()
    doc.save(buffer)
    client.post(
        "/api/v1/knowledge/documents",
        data={"title": "Monitor Manual"},
        files={
            "file": (
                "monitor.docx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    response = client.post(
        "/api/v1/analysis/questions",
        json={"question": "What should be checked for an ECG lead-off warning?", "top_k": 3},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["provider"] == "mock"
    assert body["risk_level"] == "unknown"
    assert body["citations"][0]["title"] == "Monitor Manual"
    assert "Demo mode" in body["answer"]


def test_duplicate_document_is_rejected(client):
    files = {"file": ("same.txt", b"same public demo content", "text/plain")}
    assert client.post("/api/v1/knowledge/documents", files=files).status_code == 202

    duplicate = client.post("/api/v1/knowledge/documents", files=files)

    assert duplicate.status_code == 409
    assert duplicate.json()["code"] == "document_duplicate"


def test_unsupported_document_format_is_rejected(client):
    response = client.post(
        "/api/v1/knowledge/documents",
        files={"file": ("archive.zip", b"not a document", "application/zip")},
    )

    assert response.status_code == 415
    assert response.json()["code"] == "unsupported_document_format"


def test_chinese_query_has_nonzero_lexical_and_dense_scores(client):
    document = "设备重复告警时，先检查连接线和供电状态。".encode()
    client.post(
        "/api/v1/knowledge/documents",
        files={"file": ("guide.md", document, "text/markdown")},
    )

    response = client.post(
        "/api/v1/knowledge/search", json={"query": "重复告警 检查", "top_k": 1}
    )

    assert response.status_code == 200
    scores = response.json()[0]["scores"]
    assert scores["keyword_score"] > 0
    assert scores["dense_score"] > 0
